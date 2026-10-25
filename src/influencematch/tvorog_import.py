"""Checksum-verified adapter for the CC BY educational Reels dataset."""
import argparse
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import re
from statistics import mean, median
from urllib.parse import urlsplit, urlunsplit

import requests
from openpyxl import load_workbook

from .provenance import sha256_file

RECORD_ID = "23301344"
SOURCE_URL = "https://doi.org/10.5281/zenodo.23301344"
DOWNLOAD_URL = "https://zenodo.org/api/records/23301344/files/Reels%20analysis.xlsx/content"
EXPECTED_MD5 = "b86000d0de934c01774bc9e164ea33d2"
SOURCE_CACHE = Path("data/source/zenodo_23301344/Reels analysis.xlsx")
REQUIRED_COLUMNS = {
    "Video ID", "URL", "Author", "Followers", "Publication date", "Analysis date",
    "Topic", "Views", "Comments", "Education value",
    "Attention retention potential", "Content execution quality",
}


def _source_file() -> Path:
    if not SOURCE_CACHE.is_file():
        SOURCE_CACHE.parent.mkdir(parents=True, exist_ok=True)
        response = requests.get(DOWNLOAD_URL, timeout=60)
        response.raise_for_status()
        temporary = SOURCE_CACHE.with_suffix(".xlsx.tmp")
        temporary.write_bytes(response.content)
        temporary.replace(SOURCE_CACHE)
    digest = hashlib.md5(SOURCE_CACHE.read_bytes()).hexdigest()
    if digest != EXPECTED_MD5:
        raise ValueError(f"Zenodo workbook checksum mismatch: expected {EXPECTED_MD5}, got {digest}")
    return SOURCE_CACHE


def _number(value, field: str) -> float:
    if isinstance(value, bool) or value is None:
        raise ValueError(f"Missing numeric value in {field}")
    result = float(value)
    if not math.isfinite(result):
        raise ValueError(f"Invalid numeric value in {field}")
    return result


def _source_count(value, field: str) -> tuple[int, str, bool]:
    source_value = str(value).strip()
    if isinstance(value, (int, float)):
        return int(value), source_value, False
    compact = source_value.replace(",", "").upper()
    match = re.fullmatch(r"([0-9]+(?:\.[0-9]+)?)([KM]?)", compact)
    if not match:
        raise ValueError(f"Unrecognized {field} value: {source_value!r}")
    number = float(match.group(1))
    multiplier = {"": 1, "K": 1_000, "M": 1_000_000}[match.group(2)]
    estimate = int(number * multiplier)
    return estimate, source_value, bool(match.group(2))


def _date(value, field: str) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    try:
        return date.fromisoformat(str(value)[:10]).isoformat()
    except ValueError as exc:
        raise ValueError(f"Invalid date in {field}: {value!r}") from exc


def _clean_url(value) -> str:
    parts = urlsplit(str(value).strip())
    if parts.scheme != "https" or parts.netloc.casefold() not in {"instagram.com", "www.instagram.com"}:
        raise ValueError("Workbook contains an unexpected Reel URL")
    return urlunsplit((parts.scheme, parts.netloc, parts.path, "", ""))


def import_tvorog(destination: Path = Path("data/active")) -> dict:
    source = _source_file()
    workbook = load_workbook(source, read_only=True, data_only=True)
    if "Analysis" not in workbook.sheetnames:
        raise ValueError("Zenodo workbook is missing its Analysis sheet")
    sheet = workbook["Analysis"]
    sheet.reset_dimensions()
    rows = sheet.iter_rows(values_only=True)
    next(rows, None)  # Category headings
    header = next(rows, None)
    if header is None:
        raise ValueError("Zenodo Analysis sheet has no column headers")
    columns = {str(name).strip(): index for index, name in enumerate(header) if name is not None}
    missing = REQUIRED_COLUMNS - columns.keys()
    if missing:
        raise ValueError(f"Zenodo Analysis sheet is missing columns: {', '.join(sorted(missing))}")

    by_author: dict[str, list[dict]] = defaultdict(list)
    reel_records = []
    seen_video_ids = set()
    for line_number, row in enumerate(rows, 3):
        if row[columns["Video ID"]] is None:
            continue
        author = str(row[columns["Author"]] or "").strip().lstrip("@")
        video_id = str(row[columns["Video ID"]]).removesuffix(".0")
        if not author or not video_id:
            raise ValueError(f"Missing author or video ID on Analysis row {line_number}")
        if video_id in seen_video_ids:
            raise ValueError(f"Duplicate video ID on Analysis row {line_number}")
        seen_video_ids.add(video_id)
        followers, followers_source_value, approximate_followers = _source_count(row[columns["Followers"]], "Followers")
        views, views_source_value, approximate_views = _source_count(row[columns["Views"]], "Views")
        comments, comments_source_value, approximate_comments = _source_count(row[columns["Comments"]], "Comments")
        if followers < 0 or views < 0 or comments < 0:
            raise ValueError(f"Negative follower, view, or comment count on Analysis row {line_number}")
        creator_id = f"zenodo-{RECORD_ID}:{author.casefold()}"
        reel = {
            "video_id": video_id,
            "creator_id": creator_id,
            "author": author,
            "url": _clean_url(row[columns["URL"]]),
            "topic": str(row[columns["Topic"]] or "").strip(),
            "followers": followers,
            "followers_source_value": followers_source_value,
            "approximate_followers": approximate_followers,
            "views": views,
            "views_source_value": views_source_value,
            "views_are_approximate": approximate_views,
            "comments": comments,
            "comments_source_value": comments_source_value,
            "comments_are_approximate": approximate_comments,
            "publication_date": _date(row[columns["Publication date"]], "Publication date"),
            "analysis_date": _date(row[columns["Analysis date"]], "Analysis date"),
            "education_value": _number(row[columns["Education value"]], "Education value"),
            "attention_retention_potential": _number(row[columns["Attention retention potential"]], "Attention retention potential"),
            "content_execution_quality": _number(row[columns["Content execution quality"]], "Content execution quality"),
        }
        by_author[author].append(reel)
        reel_records.append(reel)
    workbook.close()

    if len(seen_video_ids) != 300 or len(by_author) != 21:
        raise ValueError(f"Unexpected published dataset shape: {len(seen_video_ids)} Reels, {len(by_author)} accounts")

    profiles = []
    for author, reels in sorted(by_author.items(), key=lambda item: item[0].casefold()):
        topics = Counter(reel["topic"] for reel in reels if reel["topic"])
        follower_values = {reel["followers"] for reel in reels}
        if len(follower_values) != 1:
            raise ValueError(f"Follower count changes within workbook records for account {author!r}")
        source_labels = {reel["followers_source_value"] for reel in reels}
        profiles.append({
            "creator_id": f"zenodo-{RECORD_ID}:{author.casefold()}",
            "handle": author,
            "platform": "instagram",
            "followers": next(iter(follower_values)),
            "bio": "",
            "primary_topic": " ".join(topic for topic, count in topics.most_common() for _ in range(count)),
            "country": "",
            "language": "",
            "observed_at": max(reel["analysis_date"] for reel in reels),
            "source_metrics": {
                "reels_observed": len(reels),
                "mean_views": round(mean(reel["views"] for reel in reels), 2),
                "median_views": median(reel["views"] for reel in reels),
                "mean_comments": round(mean(reel["comments"] for reel in reels), 2),
                "median_comments": median(reel["comments"] for reel in reels),
                "mean_author_coded_education_value": round(mean(reel["education_value"] for reel in reels), 3),
                "mean_author_coded_attention_retention_potential": round(mean(reel["attention_retention_potential"] for reel in reels), 3),
                "mean_author_coded_content_execution_quality": round(mean(reel["content_execution_quality"] for reel in reels), 3),
                "topic_counts": dict(topics),
                "reel_urls": [reel["url"] for reel in reels],
                "followers_source_values": sorted(source_labels),
                "followers_are_approximate": any(reel["approximate_followers"] for reel in reels),
                "views_are_approximate": any(reel["views_are_approximate"] for reel in reels),
                "comments_are_approximate": any(reel["comments_are_approximate"] for reel in reels),
                "observation_window": "2026-09-02 to 2026-10-01",
            },
            "source_metric_note": (
                "Views and comments were recorded once per Reel during the source coding window; they are not live. "
                "No likes are supplied, so no engagement rate is computed. Three quality fields are one researcher's "
                "ratings, not campaign outcomes. Abbreviated follower/view/comment counts are converted to rounded estimates."
            ),
        })

    destination.mkdir(parents=True, exist_ok=True)
    output = destination / "creators.jsonl"
    temporary_output = output.with_suffix(".jsonl.tmp")
    with temporary_output.open("w", encoding="utf-8", newline="\n") as stream:
        for profile in profiles:
            stream.write(json.dumps(profile, ensure_ascii=False, allow_nan=False) + "\n")
    temporary_output.replace(output)
    reels_output = destination / "reels.jsonl"
    temporary_reels = reels_output.with_suffix(".jsonl.tmp")
    with temporary_reels.open("w", encoding="utf-8", newline="\n") as stream:
        for reel in reel_records:
            stream.write(json.dumps(reel, ensure_ascii=False, allow_nan=False) + "\n")
    temporary_reels.replace(reels_output)
    manifest = {
        "dataset_name": "Educational Instagram Reels: Coded Dataset of 300 Reels from 21 Accounts",
        "dataset_key": "zenodo_23301344_educational_reels",
        "source_url": SOURCE_URL,
        "license_id": "CC-BY-4.0",
        "source_collection_date": "2026-09-02 to 2026-10-01",
        "source_collection_window": {"start": "2026-09-02", "end": "2026-10-01"},
        "retrieved_at_utc": datetime.now(timezone.utc).isoformat(),
        "source_file": source.name,
        "source_file_size_bytes": source.stat().st_size,
        "source_file_md5": EXPECTED_MD5,
        "source_file_sha256": sha256_file(source),
        "source_reel_count": len(seen_video_ids),
        "reels_file": reels_output.name,
        "reels_file_sha256": sha256_file(reels_output),
        "reel_count": len(reel_records),
        "imported_file": output.name,
        "imported_file_sha256": sha256_file(output),
        "record_count": len(profiles),
        "data_status": "licensed_research_data",
        "rights_attestation": "Zenodo record publishes this workbook under CC BY 4.0; attribution required.",
        "attribution": "Tvorogov, Grigorii (2026), Educational Instagram Reels: Coded Dataset of 300 Reels from 21 Accounts, Zenodo, DOI 10.5281/zenodo.23301344, CC BY 4.0.",
        "limitations": [
            "The dataset covers educational Reels from 21 public accounts, not a general influencer inventory.",
            "Metrics are historical and were recorded once per Reel during the coding period.",
            "No likes, creator country, audience geography, language, audience composition, or creator rate cards are supplied.",
            "Follower, view, and comment counts use abbreviated labels for some rows and are converted to rounded estimates.",
            "Three quality fields were rated by one researcher and are not campaign outcomes.",
        ],
    }
    manifest_path = destination / "dataset_manifest.json"
    temporary_manifest = manifest_path.with_suffix(".json.tmp")
    temporary_manifest.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    temporary_manifest.replace(manifest_path)
    return manifest


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Import the checksum-verified, CC BY educational Instagram Reels dataset")
    parser.add_argument("--destination", type=Path, default=Path("data/active"))
    args = parser.parse_args()
    print(json.dumps(import_tvorog(args.destination), indent=2, ensure_ascii=False))
